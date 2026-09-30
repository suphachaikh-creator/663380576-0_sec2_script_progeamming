# src/spreadsheet_agent.py
import openpyxl
from openpyxl.styles import Font, PatternFill, Color, Alignment
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule, IconSetRule, Rule
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.utils import get_column_letter, range_boundaries
import logging
import datetime
import json

from utils import setup_logging

logger = setup_logging(__name__)

class SpreadsheetAgent:
    def __init__(self, config):
        self.config = config
        self.workbook = None
        self.audit_log = []
        self.start_time = datetime.datetime.now()
        logger.info("Spreadsheet Agent initialized.")

    def _log_audit(self, status, message, task_type="N/A", details=None):
        log_entry = {
            "timestamp": datetime.datetime.now().isoformat(),
            "status": status,
            "task_type": task_type,
            "message": message,
            "details": details if details else {}
        }
        self.audit_log.append(log_entry)
        if status in ["ERROR", "CRITICAL"]:
            logger.error(f"AUDIT LOG - {status}: {message} | Task: {task_type}")
        else:
            logger.info(f"AUDIT LOG - {status}: {message} | Task: {task_type}")

    def _load_workbook(self):
        input_file = self.config["input_file"]
        try:
            self.workbook = openpyxl.load_workbook(input_file)
            self._log_audit("SUCCESS", "Input workbook loaded successfully.", "load_workbook")
            return True
        except FileNotFoundError:
            self._log_audit("ERROR", f"Input file not found: {input_file}", "load_workbook")
            return False
        except Exception as e:
            self._log_audit("ERROR", f"Failed to load workbook: {e}", "load_workbook")
            return False

    def _get_sheet(self, sheet_name):
        if sheet_name in self.workbook.sheetnames:
            return self.workbook[sheet_name]
        else:
            logger.warning(f"Sheet '{sheet_name}' not found. Creating new sheet.")
            return self.workbook.create_sheet(sheet_name)

    def _process_copy_data_task(self, task):
        src_sheet_name = task.get("source_sheet")
        dest_sheet_name = task.get("dest_sheet")
        if not src_sheet_name or not dest_sheet_name:
            self._log_audit("ERROR", "Missing source/dest sheet", "copy_data", task)
            return
        src_sheet = self._get_sheet(src_sheet_name)
        dest_sheet = self._get_sheet(dest_sheet_name)
        for row in src_sheet.iter_rows():
            for cell in row:
                dest_sheet[cell.coordinate].value = cell.value
        self._log_audit("SUCCESS", f"Copied '{src_sheet_name}' to '{dest_sheet_name}'", "copy_data")

    def _process_calculate_column_task(self, task):
        sheet_name = task.get("sheet")
        target_col = task.get("target_column")
        start_row = task.get("start_row", 2)
        formula = task.get("formula")
        if not all([sheet_name, target_col, formula]):
            self._log_audit("ERROR", "Missing fields for calculate_column", "calculate_column", task)
            return
        sheet = self._get_sheet(sheet_name)
        header_name = task.get("header", "Calculated Column")
        sheet[f"{target_col}1"].value = header_name
        sheet[f"{target_col}1"].font = Font(bold=True)
        for row_num in range(start_row, sheet.max_row + 1):
            cell_coord = f"{target_col}{row_num}"
            try:
                sheet[cell_coord] = formula.replace("{row}", str(row_num))
                if task.get("number_format"):
                    sheet[cell_coord].number_format = task["number_format"]
            except Exception as e:
                self._log_audit("ERROR", f"Failed formula at {cell_coord}: {e}", "calculate_column")
        self._log_audit("SUCCESS", f"Column '{target_col}' calculated in '{sheet_name}'", "calculate_column")

    def _process_conditional_format_task(self, task):
        sheet_name = task.get("sheet")
        data_range = task.get("range")
        rule_type = task.get("rule_type")
        if not all([sheet_name, data_range, rule_type]):
            self._log_audit("ERROR", "Missing fields for conditional_format", "conditional_format", task)
            return
        sheet = self._get_sheet(sheet_name)
        try:
            if rule_type == "color_scale":
                min_color = Color(task.get("min_color", "FF0000"))
                mid_color = Color(task.get("mid_color", "FFFF00"))
                max_color = Color(task.get("max_color", "00FF00"))
                rule = ColorScaleRule(start_type='min', start_color=min_color,
                                     mid_type='percentile', mid_value=50, mid_color=mid_color,
                                     end_type='max', end_color=max_color)
                sheet.conditional_formatting.add(data_range, rule)
            elif rule_type == "expression":
                formula = task.get("formula")
                fill_color = task.get("fill_color", "FFC7CE")
                font_color = task.get("font_color", "9C0006")
                dxf = DifferentialStyle(
                    font=Font(color=font_color),
                    fill=PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
                )
                rule = Rule(type='expression', formula=[formula], dxf=dxf)
                sheet.conditional_formatting.add(data_range, rule)
            self._log_audit("SUCCESS", f"Applied '{rule_type}' to '{data_range}'", "conditional_format")
        except Exception as e:
            self._log_audit("ERROR", f"Failed conditional format: {e}", "conditional_format", task)

    def _process_chart_task(self, task):
        sheet_name = task.get("sheet")
        chart_type = task.get("chart_type")
        data_range = task.get("data_range")
        category_range = task.get("category_range")
        title = task.get("title", "Chart")
        top_left_cell = task.get("top_left_cell", "G2")
        if not all([sheet_name, chart_type, data_range]):
            self._log_audit("ERROR", "Missing fields for chart", "chart", task)
            return
        sheet = self._get_sheet(sheet_name)
        try:
            if chart_type == "bar":
                chart = BarChart()
            elif chart_type == "line":
                chart = LineChart()
            elif chart_type == "pie":
                chart = PieChart()
            else:
                self._log_audit("ERROR", f"Unsupported chart: {chart_type}", "chart", task)
                return
            min_col, min_row, max_col, max_row = range_boundaries(data_range)
            data = Reference(sheet, min_col=min_col, min_row=min_row, max_col=max_col, max_row=max_row)
            chart.add_data(data, titles_from_data=True)
            if category_range:
                mc, mr, xc, xr = range_boundaries(category_range)
                cats = Reference(sheet, min_col=mc, min_row=mr, max_col=xc, max_row=xr)
                chart.set_categories(cats)
            chart.title = title
            chart.x_axis.title = task.get("x_axis_title", "")
            chart.y_axis.title = task.get("y_axis_title", "")
            chart.legend = None
            sheet.add_chart(chart, top_left_cell)
            self._log_audit("SUCCESS", f"Chart '{title}' added", "chart")
        except Exception as e:
            self._log_audit("ERROR", f"Failed chart: {e}", "chart", task)

    def _add_audit_log_sheet(self):
        if not self.audit_log:
            return
        audit_sheet = self.workbook.create_sheet("Audit Log", 0)
        audit_sheet.append(["Timestamp", "Status", "Task Type", "Message", "Details"])
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="333333", end_color="333333", fill_type="solid")
        for col_num in range(1, 6):
            cell = audit_sheet.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
        for entry in self.audit_log:
            audit_sheet.append([
                entry["timestamp"], entry["status"], entry["task_type"],
                entry["message"], json.dumps(entry["details"])
            ])
        audit_widths = {"A": 24, "B": 12, "C": 22, "D": 55, "E": 18}
        for column_letter, width in audit_widths.items():
            audit_sheet.column_dimensions[column_letter].width = width
        for row in audit_sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=cell.column in (4, 5))
            audit_sheet.row_dimensions[row[0].row].height = 30
        logger.info("Audit log sheet created.")

    def run(self):
        if not self._load_workbook():
            self._log_audit("CRITICAL", "Agent could not start", "run")
            return
        try:
            for i, task in enumerate(self.config["tasks"]):
                task_type = task.get("type")
                logger.info(f"Executing Task {i+1}: Type='{task_type}'")
                if task_type == "copy_data":
                    self._process_copy_data_task(task)
                elif task_type == "calculate_column":
                    self._process_calculate_column_task(task)
                elif task_type == "conditional_format":
                    self._process_conditional_format_task(task)
                elif task_type == "chart":
                    self._process_chart_task(task)
                else:
                    self._log_audit("WARNING", f"Unsupported task: {task_type}", "task_execution")
            self._add_audit_log_sheet()
            # Fit source and report columns to their contents so headers are readable.
            for sheet in self.workbook.worksheets:
                if sheet.title == "Audit Log":
                    continue
                for column_cells in sheet.columns:
                    max_length = max(
                        (len(str(cell.value)) for cell in column_cells if cell.value is not None),
                        default=0,
                    )
                    column_letter = get_column_letter(column_cells[0].column)
                    sheet.column_dimensions[column_letter].width = min(max_length + 2, 40)
            output_file = self.config["output_file"]
            self.workbook.save(output_file)
            self._log_audit("SUCCESS", f"Workbook saved to '{output_file}'", "save_workbook")
            logger.info(f"✅ Final workbook saved to: {output_file}")
        except Exception as e:
            self._log_audit("CRITICAL", f"Unhandled error: {e}", "run")
            logger.critical(f"Error: {e}", exc_info=True)
        finally:
            duration = (datetime.datetime.now() - self.start_time).total_seconds()
            logger.info(f"⏱️  Agent finished in {duration:.2f} seconds.")
