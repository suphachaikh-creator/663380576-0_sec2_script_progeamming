# agentic-document-automator/src/word_tasks.py
import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os
from utils import setup_logging  # ✅ แก้จาก .utils เป็น utils

logger = setup_logging(__name__)

class WordTasks:
    """Handles advanced Word document operations for the Document Agent."""
    
    def __init__(self):
        logger.info("WordTasks initialized.")

    def _replace_placeholder_in_paragraph(self, paragraph, placeholder, value):
        """Replaces a single placeholder in a paragraph's runs."""
        token = "{{" + str(placeholder).strip("{}") + "}}"
        if token in paragraph.text:
            inline = paragraph.runs
            for item in inline:
                if token in item.text:
                    item.text = item.text.replace(token, str(value))
                    logger.debug(f"Replaced '{token}' with '{value}' in paragraph.")
            return True  # Placeholder found and replaced
        return False  # Placeholder not found in this paragraph

    def generate_report_from_template(self, template_path, output_path, data_placeholders, table_data=None):
        """
        Generates a Word document from a template, filling placeholders and optional table data.
        `data_placeholders` is a dictionary where keys are placeholders (e.g., '{{REPORT_TITLE}}') 
        and values are the replacement strings.
        `table_data` is a list of lists, where the first sub-list is the header row.
        """
        if not os.path.exists(template_path):
            logger.error(f"Word template not found: {template_path}")
            return False
        
        try:
            document = docx.Document(template_path)
            
            # Replace placeholders in paragraphs
            for paragraph in document.paragraphs:
                for placeholder, value in data_placeholders.items():
                    self._replace_placeholder_in_paragraph(paragraph, placeholder, value)
            
            # Replace placeholders in table cells
            for table in document.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            for placeholder, value in data_placeholders.items():
                                self._replace_placeholder_in_paragraph(paragraph, placeholder, value)
            
            # Add dynamic table if provided
            if table_data:
                logger.info("Adding dynamic table to the document.")
                rows = len(table_data)
                cols = len(table_data[0]) if table_data else 0
                
                # Find a specific placeholder to insert table, or append to end
                table_anchor_found = False
                for i, paragraph in enumerate(document.paragraphs):
                    if "{{TABLE_PLACEHOLDER}}" in paragraph.text:
                        p = paragraph._element
                        parent = p.getparent()
                        insert_at = parent.index(p)
                        parent.remove(p)
                        
                        # ✅ แก้: ใช้ document.add_table() แทน OxmlElement แบบ manual
                        table = document.add_table(rows=0, cols=cols)
                        table_element = table._tbl
                        table_element.getparent().remove(table_element)
                        parent.insert(insert_at, table_element)
                        
                        # ✅ แก้: ห่อ try-except เพื่อป้องกัน error เมื่อ style ไม่มีอยู่จริง
                        try:
                            table.style = 'Table Grid'
                        except Exception:
                            logger.warning("Style 'Table Grid' not available, using default table style.")
                        
                        # Add rows and cells from data
                        for r_idx, row_data in enumerate(table_data):
                            row = table.add_row()
                            for c_idx, cell_data in enumerate(row_data):
                                cell = row.cells[c_idx]
                                cell.text = str(cell_data)
                                # ✅ แก้: ตรวจสอบว่ามี runs ก่อนถึงจะ bold ได้
                                if r_idx == 0 and cell.paragraphs[0].runs:
                                    cell.paragraphs[0].runs[0].bold = True
                        
                        table_anchor_found = True
                        break  # Only replace first placeholder
                
                if not table_anchor_found:
                    logger.warning("Table placeholder '{{TABLE_PLACEHOLDER}}' not found in template. Appending table to end.")
                    table = document.add_table(rows=0, cols=cols)
                    
                    # ✅ แก้: ห่อ try-except เช่นกัน
                    try:
                        table.style = 'Table Grid'
                    except Exception:
                        logger.warning("Style 'Table Grid' not available, using default table style.")
                    
                    for r_idx, row_data in enumerate(table_data):
                        row = table.add_row()
                        for c_idx, cell_data in enumerate(row_data):
                            cell = row.cells[c_idx]
                            cell.text = str(cell_data)
                            # ✅ แก้: ตรวจสอบว่ามี runs ก่อนถึงจะ bold ได้
                            if r_idx == 0 and cell.paragraphs[0].runs:
                                cell.paragraphs[0].runs[0].bold = True
            
            # Save the new document
            document.save(output_path)
            logger.info(f"Report generated from template and saved to '{output_path}'.")
            return True
            
        except Exception as e:
            logger.error(f"Error generating report from template: {e}")
            return False

    def extract_text_from_word(self, docx_path):
        """Extracts all text from a Word document."""
        if not os.path.exists(docx_path):
            logger.error(f"Word document not found: {docx_path}")
            return None
        
        try:
            document = docx.Document(docx_path)
            full_text = []
            
            for para in document.paragraphs:
                full_text.append(para.text)
            
            # Also extract text from tables
            for table in document.tables:
                for row in table.rows:
                    row_text = []
                    for cell in row.cells:
                        row_text.append(cell.text)
                    full_text.append("\t".join(row_text))  # Tab-separated cell values
            
            logger.info(f"Text extracted from Word document: {docx_path}")
            return "\n".join(full_text)
            
        except Exception as e:
            logger.error(f"Error extracting text from Word document '{docx_path}': {e}")
            return None
